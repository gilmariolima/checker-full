import asyncio
import io
import unittest
from datetime import datetime

from openpyxl import Workbook

from servidor import processar_excel


class ExcelSetorTests(unittest.TestCase):
    def ler_agentes(self, setores, mesclar=True, data="07/09/2026"):
        workbook = Workbook()
        sheet = workbook.active
        for indice, setor in enumerate(setores):
            row = indice * 4 + 1
            sheet.cell(row, 1, "AGENTE: ANDRE LUCAS")
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            sheet.cell(row, 3, data)
            sheet.cell(row, 4 if mesclar else 5, setor)
            if mesclar:
                sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
            for coluna, valor in enumerate(["Cliente", "11:20", 47.38, 0, 47.38], 1):
                sheet.cell(row + 1, coluna, valor)
        buffer = io.BytesIO()
        workbook.save(buffer)
        resultado = asyncio.run(processar_excel(buffer.getvalue()))
        return [item["agente"] for item in resultado["tabela"]]

    def test_setores_mesclados_por_agente(self):
        setores = ["Vale Viagens", "Suporte Online", "Top Viagens", "Canoa"]
        self.assertEqual(
            self.ler_agentes(setores),
            [f"ANDRE LUCAS - {setor.upper()}" for setor in setores],
        )

    def test_setor_na_ultima_coluna(self):
        self.assertEqual(self.ler_agentes(["Canoa"], mesclar=False), ["ANDRE LUCAS - CANOA"])

    def test_setor_ausente_nao_vira_nan_nem_reutiliza_anterior(self):
        for data in ["07/09/2026", datetime(2026, 9, 7)]:
            with self.subTest(data=data):
                self.assertEqual(
                    self.ler_agentes(["Suporte Online", None], data=data),
                    ["ANDRE LUCAS - SUPORTE ONLINE", "ANDRE LUCAS"],
                )


if __name__ == "__main__":
    unittest.main()
