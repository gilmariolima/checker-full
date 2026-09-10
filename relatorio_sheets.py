"""Regras dos relatórios Geral/Agência, sem acesso à rede ou persistência."""

import io
import math
import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


AGENTES_AGENCIA = (
    "David Elias", "Ermesson Lima", "Livia Oliveira",
    "Mathias Lima", "Noe Lemos", "Vitoria Regia",
)
EPOCA_SHEETS = datetime(1899, 12, 30)


def normalizar_nome(texto):
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[_\s]+", " ", texto).strip().upper()


def data_celula(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            return (EPOCA_SHEETS + timedelta(days=valor)).date()
        except (ValueError, OverflowError):
            return None
    return None


def hora_celula(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%H:%M")
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if not math.isfinite(valor):
            return ""
        minutos = round((valor % 1) * 24 * 60) % (24 * 60)
        return f"{minutos // 60:02}:{minutos % 60:02}"
    return str(valor if valor is not None else "").strip()


def moeda(valor):
    try:
        numero = Decimal(str(valor or 0))
        return numero if numero.is_finite() else Decimal(0)
    except InvalidOperation:
        return Decimal(0)


def montar_blocos(abas, data_iso, tipo="geral"):
    """Abas: [{titulo, linhas}], linhas A3:G com UNFORMATTED_VALUE/SERIAL_NUMBER."""
    dia = date.fromisoformat(data_iso)
    if tipo not in ("geral", "agencia"):
        raise ValueError("Escolha Geral ou Agência.")
    permitidos = {normalizar_nome(nome) for nome in AGENTES_AGENCIA}
    blocos = []
    for aba in abas:
        titulo = aba["titulo"]
        if not titulo.upper().startswith("BASE_"):
            continue
        nome = titulo[5:].replace("_", " ").strip().title()
        if tipo == "agencia" and normalizar_nome(nome) not in permitidos:
            continue
        do_dia = [linha for linha in aba["linhas"] if linha and data_celula(linha[0]) == dia]
        if not do_dia:
            continue
        setor = str(do_dia[0][6] or "").strip() if len(do_dia[0]) > 6 else ""
        linhas = []
        for origem in do_dia:
            linha = list(origem) + [None] * max(0, 7 - len(origem))
            cliente = re.sub(r"^nome\b\s*", "", str(linha[1] or "").strip(), flags=re.I)
            pix, taxa = moeda(linha[3]), moeda(linha[4])
            linhas.append([cliente, hora_celula(linha[2]), float(pix), float(taxa), float(pix + taxa)])
        linhas.sort(key=lambda linha: linha[1])
        blocos.append({"agente": nome.upper(), "setor": setor, "linhas": linhas})
    blocos.sort(key=lambda bloco: normalizar_nome(bloco["agente"]))
    return blocos


def gerar_excel(blocos, data_iso):
    """Gera em memória o layout compatível com processar_excel."""
    dia = date.fromisoformat(data_iso).strftime("%d/%m/%Y")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Geral"
    for coluna, largura in zip("ABCDE", (32, 14, 18, 18, 20)):
        sheet.column_dimensions[coluna].width = largura

    def escrever(valores):
        sheet.append(valores)
        for cell in sheet[sheet.max_row]:
            # Nomes/setores são texto, mesmo quando começam com '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"

    def cabecalho(linha, cor):
        for cell in sheet[linha]:
            cell.fill = PatternFill("solid", fgColor=cor)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")

    escrever([dia])
    sheet.merge_cells("A1:E1")
    cabecalho(1, "1C4587")
    escrever([None] * 5)
    for bloco in blocos:
        escrever(["AGENTE: " + bloco["agente"], None, dia, bloco["setor"], None])
        linha = sheet.max_row
        sheet.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        sheet.merge_cells(start_row=linha, start_column=4, end_row=linha, end_column=5)
        cabecalho(linha, "990000")
        escrever(["NOME", "HORA", "PIX", "TAXA", "PIX TOTAL"])
        cabecalho(sheet.max_row, "990000")
        inicio = sheet.max_row + 1
        for lancamento in bloco["linhas"]:
            escrever(lancamento)
        totais = [float(sum((moeda(linha[i]) for linha in bloco["linhas"]), Decimal(0))) for i in (2, 3, 4)]
        escrever(["TOTAL", None, *totais])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=2)
        for row in sheet.iter_rows(min_row=inicio, max_row=sheet.max_row):
            row[1].alignment = Alignment(horizontal="center")
            for cell in row[2:5]:
                cell.number_format = '"R$" #,##0.00'
        escrever([None] * 5)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
